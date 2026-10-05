"""Bring hand labels from the labeling workbook back into data/gold.csv.

The workbook (Label sheet) carries the same columns as gold.csv plus a
row-number column. Only the three label columns are taken from it; ad
text and metadata always come from gold.csv, so an accidental edit to an
ad in Excel can't silently change the data being evaluated.
"""

from __future__ import annotations

import csv
from pathlib import Path

from .data_io import load_gold
from .library_import import ALL_COLUMNS, read_gold_rows

LABEL_COLUMNS = ("gold_goal", "gold_issue", "notes")


def import_labels(xlsx_path: str | Path, gold_path: str | Path) -> dict[str, int]:
    from openpyxl import load_workbook

    ws = load_workbook(xlsx_path, read_only=True, data_only=True)["Label"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    missing = {"ad_id", *LABEL_COLUMNS} - set(header)
    if missing:
        raise ValueError(f"{xlsx_path}: Label sheet is missing columns {sorted(missing)}")
    idx = {name: header.index(name) for name in ("ad_id", *LABEL_COLUMNS)}

    labels: dict[str, dict[str, str]] = {}
    for values in rows:
        ad_id = values[idx["ad_id"]]
        if ad_id is None or str(ad_id).strip() == "":
            continue
        ad_id = str(ad_id).strip()
        labels[ad_id] = {c: ("" if values[idx[c]] is None else str(values[idx[c]]).strip()) for c in LABEL_COLUMNS}

    gold_path = Path(gold_path)
    gold_rows = read_gold_rows(gold_path)
    gold_ids = [r["ad_id"] for r in gold_rows]
    if set(labels) != set(gold_ids):
        extra, absent = sorted(set(labels) - set(gold_ids)), sorted(set(gold_ids) - set(labels))
        raise ValueError(f"workbook and {gold_path} cover different ads; only in workbook: {extra[:5]}, only in gold: {absent[:5]}")

    for row in gold_rows:
        row.update(labels[row["ad_id"]])
    tmp = gold_path.with_suffix(".tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(ALL_COLUMNS), extrasaction="ignore")
        writer.writeheader()
        for row in gold_rows:
            writer.writerow({col: row.get(col, "") for col in ALL_COLUMNS})
    try:
        _, gold = load_gold(tmp)  # raises on any off-codebook or half-filled label
    except ValueError:
        tmp.unlink()
        raise
    tmp.replace(gold_path)
    return {"ads": len(gold_ids), "labeled": len(gold)}
