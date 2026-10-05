import csv, math, sys
sys.path.insert(0, "src")
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from adclass.schema import GOALS, ISSUES

rows = list(csv.DictReader(open("data/gold.csv", encoding="utf-8")))
N = len(rows)
F = "Arial"
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
head_font = Font(name=F, size=10, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="44546A")
input_fill = PatternFill("solid", fgColor="FFF2CC")   # pale yellow = fill these in
grey = Font(name=F, size=9, color="808080")
wrap_top = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
thin = Border(bottom=Side(style="thin", color="D9D9D9"))

wb = Workbook()

# ---------- Label sheet ----------
ws = wb.active
ws.title = "Label"
visible = ["ad_id", "ad_text", "gold_goal", "gold_issue", "notes"]
hidden = ["page_name", "source_url", "search_term", "paid_for", "started_running"]
cols = ["#"] + visible + hidden
ws.append(cols)
for c in ws[1]:
    c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
widths = {"#": 5, "ad_id": 19, "ad_text": 95, "gold_goal": 15, "gold_issue": 21, "notes": 40}
for i, name in enumerate(cols, start=1):
    letter = ws.cell(1, i).column_letter
    ws.column_dimensions[letter].width = widths.get(name, 20)
    if name in hidden:
        ws.column_dimensions[letter].hidden = True

for n, r in enumerate(rows, start=1):
    ws.append([n] + [r[c] for c in visible] + [r.get(c, "") for c in hidden])
    row = n + 1
    for c in ws[row]:
        c.font, c.alignment, c.border = font, wrap_top, thin
    ws.cell(row, 1).font = grey
    ws.cell(row, 2).font = grey
    ws.cell(row, 2).number_format = "@"  # keep long IDs as text, never scientific notation
    for col in (4, 5, 6):
        ws.cell(row, col).fill = input_fill
    lines = sum(max(1, math.ceil(len(p) / 100)) for p in r["ad_text"].split("\n"))
    ws.row_dimensions[row].height = min(409, max(18, 13 * lines + 5))

last = N + 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{last}"

# ---------- Lists sheet (dropdown sources) ----------
ls = wb.create_sheet("Lists")
ls["A1"], ls["B1"] = "goal", "issue"
for c in (ls["A1"], ls["B1"]):
    c.font = bold
for i, g in enumerate(GOALS, start=2):
    ls.cell(i, 1, g).font = font
for i, s in enumerate(ISSUES, start=2):
    ls.cell(i, 2, s).font = font
ls.sheet_state = "hidden"

def dv(col, src, label):
    v = DataValidation(type="list", formula1=src, allow_blank=True, showDropDown=False)
    v.error = f"Pick a {label} from the list. Only codebook labels are allowed."
    v.errorTitle = f"Invalid {label}"
    v.errorStyle = "stop"
    v.prompt = f"Choose the {label} (see Guide tab)."
    v.promptTitle = label
    v.showErrorMessage = True
    v.showInputMessage = True
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}{last}")

dv("D", f"=Lists!$A$2:$A${len(GOALS)+1}", "goal")
dv("E", f"=Lists!$B$2:$B${len(ISSUES)+1}", "issue")

# Labeled rows turn from yellow to white, so remaining work is easy to see.
done_fill = PatternFill("solid", fgColor="FFFFFF")
ws.conditional_formatting.add(f"D2:E{last}", FormulaRule(formula=['D2<>""'], fill=done_fill))
ws.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['E2<>""'], fill=done_fill))

# ---------- Guide sheet ----------
gs = wb.create_sheet("Guide", 0)
gs.column_dimensions["A"].width = 24
gs.column_dimensions["B"].width = 100
def put(r, a, b="", fa=font, fb=font):
    gs.cell(r, 1, a).font = fa
    gs.cell(r, 2, b).font = fb
    gs.cell(r, 1).alignment = wrap_top
    gs.cell(r, 2).alignment = wrap_top
    if isinstance(b, str) and len(b) > 100:
        gs.row_dimensions[r].height = 13 * math.ceil(len(b) / 100) + 4

put(1, "Political ad labeling", "", Font(name=F, size=14, bold=True))
put(3, "What to edit", "On the Label tab, fill in only the yellow cells: gold_goal and gold_issue (dropdowns) and notes (free text). Everything else is source data; leave it alone.", bold)
put(4, "How", "Label from the ad text alone. Apply the goal tie-breaks in order. Add a one-line note on any close call. Label everything before running the model.", bold)
put(5, "Hidden columns", "page_name, source_url, search_term, paid_for, and started_running are hidden on purpose, so labels come from the words only (the model sees only ad_text). Leave them in the file.", bold)
put(6, "When done", "Save as .xlsx and send it back. It is converted to data/gold.csv and validated.", bold)

put(8, "Example (not a real ad)", "", bold)
put(9, "ad_text", "Polls open 7am-7pm on Nov 3. Find your polling place at example.org.")
put(10, "gold_goal", "mobilization")
put(11, "gold_issue", "other")
put(12, "notes", "rule 2: practical voting info; no policy topic")
for r in (10, 11, 12):
    gs.cell(r, 2).fill = input_fill

put(14, "GOAL", "What the ad primarily asks of the viewer", bold, bold)
goal_defs = [
    ("persuasion", "Change or reinforce what the viewer thinks about a candidate, party, or issue (contrast, attack, biography)."),
    ("mobilization", "A civic action other than money: register, request/return a ballot, find a polling place, learn early-voting dates, make a plan to vote (no named candidate), volunteer, attend, sign a petition or pledge."),
    ("fundraising", "Asks for money, even while arguing a position."),
    ("other", "None of the above (lead generation, merchandise, service announcements, thank-yous)."),
    ("Tie-break 1", "Any explicit request for money -> fundraising."),
    ("Tie-break 2", "Practical voting info, or a petition / pledge / volunteer ask -> mobilization, even if the ad also argues a position."),
    ("Tie-break 3", "Urging a vote for a named candidate or ballot measure is persuasion, even with a date, unless the ad also gives practical voting info. A get-out-the-vote ask naming no candidate is mobilization."),
]
r = 15
for a, b in goal_defs:
    put(r, a, b, bold if a.startswith("Tie") else font); r += 1

r += 1
put(r, "ISSUE", "The main policy topic: whichever gets the most words", bold, bold); r += 1
issue_defs = [
    ("economy", "Jobs, wages, prices, inflation, taxes, housing costs, trade, budget"),
    ("healthcare", "Insurance, Medicare/Medicaid, drug prices, ACA, pre-existing conditions"),
    ("abortion", "Abortion access or restrictions, reproductive rights, IVF, contraception"),
    ("immigration", "Border, deportation, asylum, legal immigration"),
    ("democracy_voting", "Voting rights, election integrity, ballot access, courts as institutions, threats to democracy"),
    ("candidate_character", "Mainly about a person's honesty, scandal, competence, or biography rather than a policy"),
    ("public_safety", "Crime, policing, guns, drugs"),
    ("other", "Anything else (climate, education, veterans, foreign policy), or no identifiable topic"),
    ("Tie-break 1", "Two issues present -> the one with more of the ad's words."),
    ("Tie-break 2", "Mobilization ad with no topic -> other."),
    ("Tie-break 3", "Fundraising ad -> the issue the appeal is built around, else other."),
]
for a, b in issue_defs:
    put(r, a, b, bold if a.startswith("Tie") else font); r += 1

r += 1
put(r, "PROGRESS", "", bold); r += 1
put(r, "Ads", N); gs.cell(r, 2).alignment = Alignment(horizontal="left"); total_row = r; r += 1
gs.cell(r, 1, "Goal labeled").font = font
gs.cell(r, 2, f"=COUNTA(Label!D2:D{last})").font = font; gs.cell(r, 2).alignment = Alignment(horizontal="left"); r += 1
gs.cell(r, 1, "Issue labeled").font = font
gs.cell(r, 2, f"=COUNTA(Label!E2:E{last})").font = font; gs.cell(r, 2).alignment = Alignment(horizontal="left"); r += 1
gs.cell(r, 1, "Remaining").font = bold
gs.cell(r, 2, f"=B{total_row}*2-B{total_row+1}-B{total_row+2}").font = bold; gs.cell(r, 2).alignment = Alignment(horizontal="left")
gs.cell(r, 3, "label cells still empty (2 per ad)").font = grey; r += 2

put(r, "Counts per label", "Watch for labels with fewer than about 5 ads; their scores will be unreliable.", bold, grey); r += 1
for g in GOALS:
    gs.cell(r, 1, f"goal: {g}").font = font
    gs.cell(r, 2, f'=COUNTIF(Label!D2:D{last},"{g}")').font = font; gs.cell(r, 2).alignment = Alignment(horizontal="left"); r += 1
for s in ISSUES:
    gs.cell(r, 1, f"issue: {s}").font = font
    gs.cell(r, 2, f'=COUNTIF(Label!E2:E{last},"{s}")').font = font; gs.cell(r, 2).alignment = Alignment(horizontal="left"); r += 1

for sheet in (ws, gs):
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:1"

wb.active = 1  # open on the Label tab
out = "gold_labeling.xlsx"
wb.save(out)
print("saved", out, N, "ads")
